import io
from django.utils import timezone
from django.http import HttpResponse

def generate_tasks_pdf(queryset, filename="csg_report.pdf"):
    """
    Generate PDF report for tasks.
    """
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Table, TableStyle, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from core.query_utils import get_report_counts

    buffer = io.BytesIO()
    
    # 1. Setup Document Template with Header/Footer
    doc = BaseDocTemplate(buffer, pagesize=landscape(A4), leftMargin=0.5*inch, rightMargin=0.5*inch, topMargin=1*inch, bottomMargin=0.8*inch)
    
    def header_footer(canvas, doc):
        canvas.saveState()
        # Header
        canvas.setFont('Helvetica-Bold', 10)
        canvas.setFillColor(colors.HexColor('#1e3a5f'))
        canvas.drawString(0.5*inch, landscape(A4)[1] - 0.5*inch, "CSG Task Management System")
        canvas.setFont('Helvetica', 9)
        canvas.setFillColor(colors.HexColor('#666666'))
        canvas.drawRightString(landscape(A4)[0] - 0.5*inch, landscape(A4)[1] - 0.5*inch, f"Generated: {timezone.now().strftime('%b %d, %Y %H:%M')}")
        
        # Footer
        canvas.setFont('Helvetica', 9)
        canvas.setFillColor(colors.HexColor('#999999'))
        canvas.drawString(0.5*inch, 0.4*inch, "Confidential - For internal use only")
        canvas.drawRightString(landscape(A4)[0] - 0.5*inch, 0.4*inch, f"Page {doc.page}")
        canvas.restoreState()

    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='normal')
    template = PageTemplate(id='Report', frames=frame, onPage=header_footer)
    doc.addPageTemplates([template])

    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=22, spaceAfter=10, textColor=colors.HexColor('#1e3a5f'), alignment=0)
    story.append(Paragraph('Task Management Report', title_style))
    story.append(HRFlowable(width='100%', thickness=2, color=colors.HexColor('#1e3a5f'), spaceAfter=15))

    # 2. Summary Analytics
    today = timezone.now().date()
    stats = get_report_counts(queryset, today)
    
    summary_data = [
        [
            Paragraph(f"<b>Total Tasks</b><br/><font size=14>{stats['total']}</font>", styles['Normal']),
            Paragraph(f"<b>Completed</b><br/><font size=14 color='#10B981'>{stats['completed']}</font>", styles['Normal']),
            Paragraph(f"<b>In Progress</b><br/><font size=14 color='#3B82F6'>{stats['in_progress']}</font>", styles['Normal']),
            Paragraph(f"<b>Overdue</b><br/><font size=14 color='#EF4444'>{stats['overdue']}</font>", styles['Normal']),
        ]
    ]
    summary_table = Table(summary_data, colWidths=[2.5*inch]*4)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 20))

    # 3. Task Data Table
    data = [['Task No.', 'Title', 'Status', 'Priority', 'Assigned To', 'Due Date', 'Progress']]
    
    if stats['total'] > 500:
        task_iter = queryset.iterator(chunk_size=200)
    else:
        task_iter = queryset
        
    for t in task_iter:
        officers_list = getattr(t, 'sorted_assigned_officers', t.assigned_officers.all())
        if officers_list:
            officers = ', '.join([f"{o.get_full_name() or o.username} ({getattr(o, 'position_initials', '')})" for o in officers_list])
        else:
            officers = 'Unassigned'
        data.append([
            t.task_number,
            Paragraph(t.title[:65], styles['Normal']),
            t.get_status_display(),
            t.get_priority_display(),
            Paragraph(officers[:40], styles['Normal']),
            str(t.due_date) if t.due_date else 'N/A',
            f'{t.progress}%'
        ])

    col_widths = [1.2*inch, 3.2*inch, 1.2*inch, 1.0*inch, 2.0*inch, 1.1*inch, 1.0*inch]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a5f')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    
    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    response['X-Frame-Options'] = 'SAMEORIGIN'
    return response


def generate_tasks_excel(queryset, filename="csg_report.xlsx"):
    """
    Generate Excel report for tasks.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, NamedStyle
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'CSG Tasks Report'

    # 1. Metadata Header
    ws.merge_cells('A1:I1')
    title_cell = ws.cell(row=1, column=1, value='CSG Task Management Report')
    title_cell.font = Font(bold=True, size=16, color='1e3a5f')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    task_count = queryset.count()
    ws.merge_cells('A2:I2')
    meta_cell = ws.cell(row=2, column=1, value=f'Generated: {timezone.now().strftime("%B %d, %Y %I:%M %p")} | Total Tasks: {task_count}')
    meta_cell.font = Font(size=11, color='666666')
    meta_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Leave a blank row
    ws.append([])

    # 2. Table Headers (Row 4)
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='1e3a5f', end_color='1e3a5f', fill_type='solid')
    alt_fill = PatternFill(start_color='EBF3FB', end_color='EBF3FB', fill_type='solid')

    headers = ['Task Number', 'Title', 'Status', 'Priority', 'Assigned Officers', 'Due Date', 'Completion Date', 'Progress (%)', 'Created By']
    ws.append(headers)
    
    header_row = 4
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Add Auto-Filtering
    ws.auto_filter.ref = f"A{header_row}:I{header_row}"

    # Set up date styles
    date_style = NamedStyle(name='date_style', number_format='YYYY-MM-DD')
    if 'date_style' not in wb.named_styles:
        wb.add_named_style(date_style)

    if task_count > 500:
        task_iter = queryset.iterator(chunk_size=200)
    else:
        task_iter = queryset

    # Define Conditional Format colors
    colors = {
        'completed': PatternFill(start_color='dcfce7', end_color='dcfce7', fill_type='solid'),
        'overdue': PatternFill(start_color='fee2e2', end_color='fee2e2', fill_type='solid'),
        'in_progress': PatternFill(start_color='e0f2fe', end_color='e0f2fe', fill_type='solid'),
    }

    start_data_row = 5
    for i, t in enumerate(task_iter, start_data_row):
        officers_list = getattr(t, 'sorted_assigned_officers', t.assigned_officers.all())
        if officers_list:
            officers = ', '.join([f"{o.get_full_name() or o.username} ({getattr(o, 'position_initials', '')})" for o in officers_list])
        else:
            officers = 'Unassigned'
        
        # Determine status color
        status_fill = None
        if t.status == 'completed':
            status_fill = colors['completed']
        elif t.is_overdue:
            status_fill = colors['overdue']
        elif t.status not in ['not_started', 'completed']:
            status_fill = colors['in_progress']

        # Append row data
        ws.append([
            t.task_number, 
            t.title, 
            t.get_status_display(), 
            t.get_priority_display(),
            officers,
            t.due_date if t.due_date else None,
            t.completion_date if t.completion_date else None,
            t.progress,
            t.created_by.get_full_name() or t.created_by.username,
        ])
        
        # Apply styles to the newly appended row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=col)
            # Alternate row background (except for status column which gets conditional color)
            if i % 2 == 0:
                cell.fill = alt_fill
            
            # Apply date format
            if col in [6, 7]:  # Due Date, Completion Date
                cell.style = 'date_style'
            
            # Apply status conditional formatting
            if col == 3 and status_fill:
                cell.fill = status_fill

    widths = [15, 45, 18, 12, 40, 15, 18, 12, 20]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    
    ws.freeze_panes = f'A{start_data_row}'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
